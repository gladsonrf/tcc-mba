"""Gera data/external/culturas_fertilizantes.xlsx (planilha de culturas, calendário e fertilizantes).

Obs.: as fórmulas só ganham valores ao abrir no Excel/LibreOffice."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

from config import RAW, EXTERNAL
UP = str(RAW / "comexstat") + "/"
EXTERNAL.mkdir(parents=True, exist_ok=True)
OUT = str(EXTERNAL / "culturas_fertilizantes.xlsx")

F = "Arial"
H_FILL = PatternFill("solid", fgColor="1F3A5F")
H_FONT = Font(name=F, bold=True, color="FFFFFF", size=10)
B_FONT = Font(name=F, size=10)
IN_FONT = Font(name=F, size=10, color="0000FF")
LINK_FONT = Font(name=F, size=10, color="008000")
TITLE = Font(name=F, bold=True, size=13, color="1F3A5F")
NOTE = Font(name=F, italic=True, size=9, color="555555")
thin = Side(style="thin", color="C9D1DA")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
MESES = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]


def header(ws, row, cols, widths=None):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.font, cell.fill, cell.border = H_FONT, H_FILL, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w


def style_row(ws, row, ncol, font=B_FONT):
    for j in range(1, ncol + 1):
        c = ws.cell(row=row, column=j)
        c.border, c.alignment = BORDER, WRAP
        c.font = Font(name=F, size=10, bold=c.font.bold, italic=c.font.italic, color=c.font.color)


wb = Workbook()

# ---------------------------------------------------------------- Leia-me
ws = wb.active
ws.title = "Leia-me"
ws["A1"] = "Planilha de culturas e fertilizantes – TCC Previsão do preço de importação de fertilizantes"
ws["A1"].font = TITLE
linhas = [
    "Objetivo: delimitar as culturas do agro que mais demandam fertilizantes e quando ocorre a compra (pré-plantio),",
    "conforme sugestão do orientador (Prof. Edward, reunião de 10/10/2026).",
    "",
    "Abas:",
    "  Culturas – resumo por cultura: exportações, nutrientes, fertilizantes-chave, calendário e prioridade no estudo.",
    "  Calendario – plantio (P) e colheita (C) por mês (Conab) e participação mensal do volume importado.",
    "  Fertilizantes_NCM – NCMs dos fertilizantes do estudo, volume e preço de 2025.",
    "  Sazonalidade_importacao – volume importado por mês (2015–2025) e testes de sazonalidade de preço.",
    "  Exportacoes_agro – exportações dos principais produtos do agro (ComexStat), 2023–2025.",
    "  Fontes – referências de cada informação.",
    "",
    "Legenda de cores: texto azul = dado de entrada (fonte citada); preto = fórmula; verde = vínculo com outra aba.",
    "Células com 'a confirmar' ainda precisam de checagem em fonte primária.",
    "Dados gerados pelos scripts do repositório (src/06_exportacoes_agro.py e src/01_baixar_comexstat.py).",
]
for i, t in enumerate(linhas, 3):
    ws.cell(row=i, column=1, value=t).font = B_FONT
ws.column_dimensions["A"].width = 120

# ---------------------------------------------------------------- Exportacoes_agro
ex = pd.read_csv(UP + "exportacao_agro_anual.csv", dtype={"headingCode": str})
pv = ex.pivot_table(index=["headingCode", "produto", "cultura"], columns="ano", values="fob_usd", aggfunc="sum")
pv = pv.sort_values(2025, ascending=False).reset_index()
we = wb.create_sheet("Exportacoes_agro")
we["A1"] = "Exportações dos principais produtos do agro – US$ bilhões (FOB)"
we["A1"].font = TITLE
header(we, 3, ["SH4", "Produto", "Cultura de origem", "2023", "2024", "2025", "Part. 2025 no grupo"],
       [8, 28, 22, 11, 11, 11, 14])
r0 = 4
for i, row in pv.iterrows():
    r = r0 + i
    we.cell(row=r, column=1, value=row["headingCode"])
    we.cell(row=r, column=2, value=row["produto"])
    we.cell(row=r, column=3, value=row["cultura"])
    for j, ano in enumerate([2023, 2024, 2025]):
        c = we.cell(row=r, column=4 + j, value=round(float(row[ano]) / 1e9, 3))
        c.font, c.number_format = IN_FONT, "0.0"
    style_row(we, r, 7)
    for j in range(4, 7):
        we.cell(row=r, column=j).font = IN_FONT
rN = r0 + len(pv) - 1
for j in range(4, 7):
    we.cell(row=rN + 1, column=j, value=f"=SUM({get_column_letter(j)}{r0}:{get_column_letter(j)}{rN})").number_format = "0.0"
for r in range(r0, rN + 1):
    c = we.cell(row=r, column=7, value=f"=F{r}/$F${rN + 1}")
    c.number_format = "0.0%"
    c.font = B_FONT
we.cell(row=rN + 1, column=2, value="Total dos produtos listados").font = Font(name=F, bold=True, size=10)
for j in range(1, 8):
    we.cell(row=rN + 1, column=j).border = BORDER
    if j >= 4:
        we.cell(row=rN + 1, column=j).font = Font(name=F, bold=True, size=10)
we.cell(row=rN + 3, column=1,
        value="Fonte: ComexStat/MDIC, exportações por posição SH4, valor FOB. Baixado via API em 10/10/2026 "
              "(src/06_exportacoes_agro.py).").font = NOTE

# ---------------------------------------------------------------- Fertilizantes_NCM
nc = pd.read_csv(UP + "importacao_ncm_mensal.csv", dtype={"coNcm": str})
y = nc[nc.ano == 2025].groupby("coNcm")[["kg", "fob_usd"]].sum()
first = nc.groupby("coNcm")["ano"].min()
nmonths = nc.groupby("coNcm").size()
info = [
    ("31021010", "Ureia (N > 45%)", "N", "Milho (1ª e 2ª safra), cana, café, algodão, trigo", "Modelo por produto (N)"),
    ("31022100", "Sulfato de amônio", "N + S", "Cana, café, milho; substituto da ureia em 2026", "Robustez (N)"),
    ("31042090", "Outros cloretos de potássio (principal NCM de KCl)", "K", "Soja, cana, café, algodão", "Modelo por produto (K) – foco soja"),
    ("31042010", "Cloreto de potássio, K2O ≤ 60%", "K", "Soja, cana", "Pouco volume – somar ao KCl"),
    ("31054000", "MAP (fosfato monoamônico)", "P + N", "Soja, milho, algodão", "Modelo por produto (P) – foco soja"),
    ("31053000", "DAP (fosfato diamônico)", "P + N", "Milho, algodão", "Série com muitos meses vazios – não modelar"),
]
wf = wb.create_sheet("Fertilizantes_NCM")
wf["A1"] = "Fertilizantes do estudo – NCM, volume e preço médio de importação (2025)"
wf["A1"].font = TITLE
header(wf, 3, ["NCM", "Descrição", "Nutriente", "Volume 2025 (Mt)", "Valor FOB 2025 (US$ mi)",
               "Preço médio 2025 (US$/t)", "Primeiro ano", "Meses com dado", "Culturas que mais usam", "Uso no TCC"],
       [11, 40, 10, 13, 14, 14, 10, 10, 40, 36])
for i, (ncm, desc, nut, cult, uso) in enumerate(info):
    r = 4 + i
    vals = [ncm, desc, nut, round(float(y.loc[ncm, "kg"]) / 1e9, 3), round(float(y.loc[ncm, "fob_usd"]) / 1e6, 1),
            f"=E{r}/D{r}", int(first.loc[ncm]), int(nmonths.loc[ncm]), cult, uso]
    for j, v in enumerate(vals, 1):
        wf.cell(row=r, column=j, value=v)
    style_row(wf, r, 10)
    for j in (4, 5, 7, 8):
        wf.cell(row=r, column=j).font = IN_FONT
    wf.cell(row=r, column=4).number_format = "0.00"
    wf.cell(row=r, column=5).number_format = "#,##0.0"
    wf.cell(row=r, column=6).number_format = "#,##0"
wf.cell(row=11, column=1, value="Preço médio = valor FOB ÷ volume (fórmula). Fonte: ComexStat/MDIC, importações por NCM, "
                                "jan/1997–ago/2026 (src/01_baixar_comexstat.py).").font = NOTE
wf.cell(row=12, column=1, value="Soja: fixação biológica de nitrogênio (Bradyrhizobium) supre quase todo o N; por isso "
                                "o foco da soja é K e P (Embrapa; MAPA/PNF).").font = NOTE

# ---------------------------------------------------------------- Sazonalidade_importacao
sh = pd.read_csv(UP + "importacao_sh4_mensal.csv", dtype={"headingCode": str})
s = sh[(sh.ano >= 2015) & (sh.ano <= 2025)].groupby(["headingCode", "mes"])["kg"].sum().unstack()
wsz = wb.create_sheet("Sazonalidade_importacao")
wsz["A1"] = "Volume importado por mês – soma 2015–2025 (milhões de t) e participação mensal"
wsz["A1"].font = TITLE
header(wsz, 3, ["Posição", "Grupo"] + MESES + ["Total"], [9, 26] + [8] * 12 + [10])
grupos = {"3102": "Nitrogenados", "3104": "Potássicos", "3105": "NPK, MAP, DAP (mistos)"}
for i, code in enumerate(["3102", "3104", "3105"]):
    r = 4 + i
    wsz.cell(row=r, column=1, value=code)
    wsz.cell(row=r, column=2, value=grupos[code])
    for m in range(1, 13):
        c = wsz.cell(row=r, column=2 + m, value=round(float(s.loc[code, m]) / 1e9, 3))
        c.number_format = "0.0"
    wsz.cell(row=r, column=15, value=f"=SUM(C{r}:N{r})").number_format = "0.0"
    style_row(wsz, r, 15)
    for m in range(3, 15):
        wsz.cell(row=r, column=m).font = IN_FONT
header(wsz, 8, ["Posição", "Grupo"] + MESES + ["Total"])
for i in range(3):
    r, rs = 9 + i, 4 + i
    wsz.cell(row=r, column=1, value=f"=A{rs}")
    wsz.cell(row=r, column=2, value=f"=B{rs}")
    for m in range(3, 15):
        L = get_column_letter(m)
        c = wsz.cell(row=r, column=m, value=f"={L}{rs}/$O${rs}")
        c.number_format = "0.0%"
    wsz.cell(row=r, column=15, value=f"=SUM(C{r}:N{r})").number_format = "0%"
    style_row(wsz, r, 15)
wsz.cell(row=7, column=1, value="Participação mensal (% do volume anual)").font = Font(name=F, bold=True, size=10)
# color scale on shares
from openpyxl.formatting.rule import ColorScaleRule
wsz.conditional_formatting.add("C9:N11", ColorScaleRule(start_type="min", start_color="FFFFFF",
                                                       end_type="max", end_color="5B8DC9"))
# price seasonality tests
wsz.cell(row=14, column=1, value="Teste de sazonalidade no preço (dummies mensais na variação log, 2000–2025, erros robustos HC1)").font = Font(name=F, bold=True, size=10)
header(wsz, 15, ["Série", "Fonte", "p-valor (F)", "Leitura"])
testes = [("Ureia (internacional)", "Banco Mundial", 0.143, "Sem sazonalidade significativa a 5%"),
          ("Cloreto de potássio (internacional)", "Banco Mundial", 0.918, "Sem sazonalidade"),
          ("DAP (internacional)", "Banco Mundial", 0.002, "Sazonalidade significativa"),
          ("Nitrogenados – importação BR (3102)", "ComexStat", 0.068, "Fraca (significativa só a 10%)"),
          ("Potássicos – importação BR (3104)", "ComexStat", 0.923, "Sem sazonalidade"),
          ("NPK/MAP – importação BR (3105)", "ComexStat", 0.016, "Sazonalidade significativa")]
for i, t in enumerate(testes):
    r = 16 + i
    for j, v in enumerate(t, 1):
        wsz.cell(row=r, column=j, value=v)
    wsz.cell(row=r, column=3).number_format = "0.000"
    style_row(wsz, r, 4)
    wsz.cell(row=r, column=3).font = IN_FONT
wsz.column_dimensions["A"].width = 34
wsz.column_dimensions["B"].width = 16
wsz.column_dimensions["D"].width = 12
wsz.cell(row=23, column=1, value="Leitura: o VOLUME importado é sazonal (K e NPK concentram mai–set, antes do plantio da soja; "
                                 "N concentra set–dez). O PREÇO é pouco sazonal, exceto fosfatados. Cálculo próprio a partir do "
                                 "ComexStat e da Pink Sheet (10/10/2026); refazer no notebook de EDA.").font = NOTE

# ---------------------------------------------------------------- Calendario
wc = wb.create_sheet("Calendario")
wc["A1"] = "Calendário agrícola (P = plantio, C = colheita, P/C = ambos) e janela de importação"
wc["A1"].font = TITLE
header(wc, 3, ["Cultura", "Região de referência"] + MESES + ["Fonte"], [26, 22] + [6] * 12 + [34])
cal = {
    ("Soja", "Centro-Oeste (MT, GO, MS)"): {9: "P", 10: "P", 11: "P", 12: "P", 1: "C", 2: "C", 3: "C", 4: "C"},
    ("Soja", "Sul (PR, RS, SC)"): {9: "P", 10: "P", 11: "P", 12: "P", 1: "P/C", 2: "C", 3: "C", 4: "C", 5: "C"},
    ("Milho 1ª safra", "Centro-Oeste"): {9: "P", 10: "P", 11: "P", 12: "P", 2: "C", 3: "C", 4: "C", 5: "C", 6: "C"},
    ("Milho 2ª safra (safrinha)", "Centro-Oeste"): {1: "P", 2: "P", 3: "P", 5: "C", 6: "C", 7: "C", 8: "C", 9: "C"},
    ("Algodão", "Centro-Oeste (MT, GO)"): {11: "P", 12: "P", 1: "P", 2: "P", 6: "C", 7: "C", 8: "C", 9: "C"},
    ("Trigo", "Sul (PR, RS, SC)"): {4: "P", 5: "P", 6: "P", 7: "P", 8: "C", 9: "C", 10: "C", 11: "C", 12: "C"},
    ("Cana-de-açúcar", "Centro-Sul"): {1: "P", 2: "P", 3: "P", 4: "C", 5: "C", 6: "C", 7: "C", 8: "C", 9: "C", 10: "P/C", 11: "P/C"},
    ("Café (arábica)", "Sudeste (MG, SP)"): {5: "C", 6: "C", 7: "C", 8: "C", 9: "C"},
}
fontes_cal = {"Soja": "Conab, Calendário de plantio e colheita (2022), p. 54–55",
              "Milho 1ª safra": "Conab (2022), p. 43", "Milho 2ª safra (safrinha)": "Conab (2022), p. 48",
              "Algodão": "Conab (2022), p. 12", "Trigo": "Conab (2022), p. 69",
              "Cana-de-açúcar": "Plantio: Agrolink (ano e meio jan–mar; de ano out–nov). Colheita: a confirmar (Conab cana)",
              "Café (arábica)": "Colheita: a confirmar (Conab café). Perene: adubação nas chuvas (out–mar), a confirmar"}
r = 4
for (cult, reg), meses in cal.items():
    wc.cell(row=r, column=1, value=cult)
    wc.cell(row=r, column=2, value=reg)
    for m in range(1, 13):
        wc.cell(row=r, column=2 + m, value=meses.get(m, ""))
        wc.cell(row=r, column=2 + m).alignment = Alignment(horizontal="center")
    wc.cell(row=r, column=15, value=fontes_cal[cult])
    style_row(wc, r, 15)
    for m in range(3, 15):
        wc.cell(row=r, column=m).alignment = Alignment(horizontal="center", vertical="center")
    r += 1
last = r - 1
rng = f"C4:N{last}"
wc.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"P"'], fill=PatternFill("solid", fgColor="F6C28B")))
wc.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"C"'], fill=PatternFill("solid", fgColor="A9C4E8")))
wc.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"P/C"'], fill=PatternFill("solid", fgColor="B8DDB0")))
r += 1
wc.cell(row=r, column=1, value="Participação mensal do volume importado (2015–2025)").font = Font(name=F, bold=True, size=10)
r += 1
for i in range(3):
    wc.cell(row=r + i, column=1, value=f"=Sazonalidade_importacao!B{9 + i}").font = LINK_FONT
    wc.cell(row=r + i, column=2, value="Brasil – importação").font = B_FONT
    for m in range(3, 15):
        L = get_column_letter(m)
        c = wc.cell(row=r + i, column=m, value=f"=Sazonalidade_importacao!{L}{9 + i}")
        c.number_format = "0%"
        c.font = LINK_FONT
        c.border = BORDER
wc.conditional_formatting.add(f"C{r}:N{r + 2}", ColorScaleRule(start_type="min", start_color="FFFFFF",
                                                              end_type="max", end_color="5B8DC9"))
wc.cell(row=r + 4, column=1, value="Leitura: K e NPK são importados sobretudo de mai a set, antes do plantio da soja (set–dez); "
                                   "N concentra set–dez (milho 1ª safra e cobertura, cana). A janela de compra precede o plantio.").font = NOTE

# ---------------------------------------------------------------- Culturas (main)
wcu = wb.create_sheet("Culturas", 1)
wcu["A1"] = "Culturas do agro, fertilizantes e janela de compra – base para delimitar o estudo"
wcu["A1"].font = TITLE
cols = ["Cultura", "Produtos exportados (SH4)", "Exportação 2025 (US$ bi)", "Nutrientes mais demandados",
        "Fertilizantes-chave (NCM)", "Principais regiões", "Plantio (Conab)", "Colheita (Conab)",
        "Janela provável de compra/importação", "Prioridade no TCC", "Observações"]
header(wcu, 3, cols, [20, 22, 12, 22, 26, 20, 16, 16, 28, 12, 44])
dados = [
    ("Soja", "1201, 2304, 1507", "Soja", "P e K (N via fixação biológica)", "KCl (31042090); MAP (31054000)",
     "MT, PR, GO, RS, MS", "Set–Dez", "Jan–Mai", "Mai–Set (pico de importação de K e NPK)", "Alta",
     "Maior exportação do agro. Robustez principal sugerida pelo orientador: modelos de KCl e MAP avaliados na janela pré-plantio."),
    ("Milho (1ª e 2ª safra)", "1005", "Milho", "N (principal), P, K", "Ureia (31021010); sulfato de amônio (31022100); MAP",
     "MT, PR, GO, MS", "1ª: Set–Dez · 2ª: Jan–Mar", "1ª: Fev–Jun · 2ª: Mai–Set", "N: Set–Dez (pico de importação de N)", "Alta",
     "Principal usuário de ureia. Milho 2ª safra é plantado logo após a soja."),
    ("Cana-de-açúcar", "1701, 2207", "Cana-de-açúcar", "K e N", "KCl; ureia; sulfato de amônio",
     "SP, GO, MG, MS, PR", "Jan–Mar (ano e meio); Out–Nov (de ano)", "Abr–Nov (a confirmar)", "Set–Dez e Jan–Mar (a confirmar)", "Alta",
     "Soja, milho e cana somam ~72% do consumo de fertilizantes (MAPA/PNF, 2020)."),
    ("Café", "0901", "Café", "N e K", "Ureia; KCl; sulfato de amônio", "MG, ES, SP", "Perene",
     "Mai–Set (a confirmar)", "Ago–Nov, antes das chuvas (a confirmar)", "Média",
     "2ª maior exportação do agro em 2025; cultura perene: adubação parcelada no período chuvoso."),
    ("Algodão", "5201", "Algodão", "N e K", "Ureia; KCl", "MT, BA, GO", "Nov–Fev", "Jun–Set",
     "Ago–Dez (a confirmar)", "Média", "Em MT, plantado como 2ª safra após a soja."),
    ("Trigo", "—", None, "N", "Ureia", "PR, RS", "Abr–Jul", "Ago–Dez", "Fev–Mai (a confirmar)", "Baixa",
     "Brasil é importador de trigo; área pequena frente a soja e milho."),
    ("Pecuária (carnes)", "0202, 0207, 0203", "Pecuária", "Indireto (milho e soja na ração); pastagens", "—",
     "Diversas", "—", "—", "—", "Baixa", "Não modelar diretamente: efeito passa por milho e soja."),
    ("Celulose, laranja, fumo", "4703, 2009, 2401", None, "Variado", "—", "Diversas", "—", "—", "—", "Baixa",
     "Pequena participação no consumo de fertilizantes."),
]
for i, d in enumerate(dados):
    r = 4 + i
    (cult, sh4, chave, nut, fert, reg, pl, co, jan, pri, obs) = d
    vals = [cult, sh4, None, nut, fert, reg, pl, co, jan, pri, obs]
    for j, v in enumerate(vals, 1):
        wcu.cell(row=r, column=j, value=v)
    if chave:
        f = f'=SUMIF(Exportacoes_agro!$C:$C,"{chave}",Exportacoes_agro!$F:$F)'
        if cult.startswith("Celulose"):
            f = ('=SUMIF(Exportacoes_agro!$C:$C,"Eucalipto/silvicultura",Exportacoes_agro!$F:$F)'
                 '+SUMIF(Exportacoes_agro!$C:$C,"Laranja",Exportacoes_agro!$F:$F)'
                 '+SUMIF(Exportacoes_agro!$C:$C,"Fumo",Exportacoes_agro!$F:$F)')
        wcu.cell(row=r, column=3, value=f)
    elif cult.startswith("Celulose"):
        wcu.cell(row=r, column=3, value='=SUMIF(Exportacoes_agro!$C:$C,"Eucalipto/silvicultura",Exportacoes_agro!$F:$F)'
                                        '+SUMIF(Exportacoes_agro!$C:$C,"Laranja",Exportacoes_agro!$F:$F)'
                                        '+SUMIF(Exportacoes_agro!$C:$C,"Fumo",Exportacoes_agro!$F:$F)')
    else:
        wcu.cell(row=r, column=3, value="—")
    style_row(wcu, r, 11)
    wcu.cell(row=r, column=3).font = LINK_FONT
    wcu.cell(row=r, column=3).number_format = "0.0"
    wcu.cell(row=r, column=3).alignment = Alignment(horizontal="center", vertical="top")
    wcu.cell(row=r, column=10).alignment = Alignment(horizontal="center", vertical="top")
last = 4 + len(dados) - 1
wcu.conditional_formatting.add(f"J4:J{last}", CellIsRule(operator="equal", formula=['"Alta"'], fill=PatternFill("solid", fgColor="F6C28B")))
wcu.conditional_formatting.add(f"J4:J{last}", CellIsRule(operator="equal", formula=['"Média"'], fill=PatternFill("solid", fgColor="FBE3C4")))
n = last + 2
for t in ["Proposta para o TCC: um modelo por fertilizante (ureia, KCl, MAP) e, como robustez, a soja (KCl e MAP) avaliada na janela pré-plantio.",
          "Exportação 2025 vem da aba Exportacoes_agro (fórmula). 'a confirmar' = checar em fonte primária (Conab cana/café).",
          "Para preencher: novas culturas entram como nova linha; mantenha o texto azul para dados de entrada e cite a fonte na aba Fontes."]:
    wcu.cell(row=n, column=1, value=t).font = NOTE
    n += 1
wcu.freeze_panes = "B4"

# ---------------------------------------------------------------- Fontes
wfo = wb.create_sheet("Fontes")
wfo["A1"] = "Fontes"
wfo["A1"].font = TITLE
header(wfo, 3, ["Informação", "Fonte", "Link / observação"], [36, 50, 70])
fontes = [
    ("Exportações do agro", "ComexStat/MDIC – exportações por SH4 (API)", "https://comexstat.mdic.gov.br – baixado em 10/10/2026"),
    ("Importações de fertilizantes", "ComexStat/MDIC – importações por SH4 e NCM (API)", "https://comexstat.mdic.gov.br – jan/1997 a ago/2026"),
    ("Preços internacionais", "Banco Mundial – Pink Sheet (out/2026)", "https://www.worldbank.org/en/research/commodity-markets"),
    ("Calendário de grãos e algodão", "Conab – Calendário de plantio e colheita de grãos no Brasil (2022)",
     "https://www.gov.br/conab/pt-br/acesso-a-informacao/institucional/publicacoes/arquivos-de-paginas/calendariozplantiozezcolheitazjunz2022.pdf"),
    ("Plantio da cana", "Agrolink – Plantio da cana-de-açúcar",
     "https://www.agrolink.com.br/culturas/cana-de-acucar/informacoes-da-cultura/pre-plantio-e-plantio/plantio-da-cana-de-acucar_438272.html"),
    ("Consumo por cultura (soja+milho+cana = 72%, 2020)", "MAPA – Plano Nacional de Fertilizantes, estatísticas do setor",
     "https://www.gov.br/agricultura/pt-br/assuntos/insumos-agropecuarios/insumos-agricolas/fertilizantes/plano-nacional-de-fertilizantes/estatisticas-do-setor"),
    ("Fixação biológica de N na soja", "Embrapa; Conecta Agro Brasil (09/09/2026)",
     "https://conectaagrobrasil.com.br/agricultura/consumo-de-fertilizantes-por-soja-milho-e-cana-no-brasil-passa-de-70-por-que-tres-culturas-concentram-tanto-adubo/"),
    ("Sazonalidade e choques", "Cálculo próprio (ComexStat e Pink Sheet)", "Refazer e documentar no notebook de análise exploratória"),
]
for i, f_ in enumerate(fontes):
    r = 4 + i
    for j, v in enumerate(f_, 1):
        wfo.cell(row=r, column=j, value=v)
    style_row(wfo, r, 3)

wb.move_sheet("Culturas", offset=-0)
wb.save(OUT)
print("ok", OUT)
